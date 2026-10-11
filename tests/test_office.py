import importlib.util
import tempfile
import unittest
from pathlib import Path
from marvis.runtime import Runtime
from marvis.capabilities import Capabilities


@unittest.skipUnless(all(importlib.util.find_spec(m) for m in ('pypdf','docx','openpyxl')), 'Optional office dependencies not installed')
class OfficeTests(unittest.TestCase):
    def test_real_word_excel_and_pdf_extraction(self):
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); work=root/'work';work.mkdir()
            runtime=Runtime(root/'data',offline=True)
            try:
                runtime.update_settings({'roots':[str(work)]});caps=Capabilities(runtime.store)
                doc=Document();doc.add_paragraph('本地 Word 文档测试');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='客户';table.cell(0,1).text='上海';doc.save(work/'sample.docx')
                self.assertIn('上海',caps.read(str(work/'sample.docx')))
                self.assertIn('本地 Word',caps.read(str(work/'sample.docx')))
                wb=Workbook();wb.active.append(['金额',125]);wb.active.append(['零',0]);wb.save(work/'sample.xlsx');wb.close()
                self.assertIn('125',caps.read(str(work/'sample.xlsx')))
                self.assertIn('零 | 0',caps.read(str(work/'sample.xlsx')))
                from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
                writer=PdfWriter(); page=writer.add_blank_page(width=300,height=200)
                font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
                page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
                stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 100 Td (Local PDF extraction) Tj ET')
                page[NameObject('/Contents')]=writer._add_object(stream);writer.write(str(work/'sample.pdf'))
                self.assertIn('Local PDF extraction',caps.read(str(work/'sample.pdf')))
            finally:runtime.close()
